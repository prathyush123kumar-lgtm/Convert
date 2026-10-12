import sys
import unittest
from pathlib import Path
import pypdf
import docx

# Add project root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from services.cleanup import EphemeralWorkspace
from services.validator import validate_upload, sanitize_filename, ValidationError
from services.office_to_pdf import convert_word_to_pdf, ConversionError
from services.pdf_to_word import convert_pdf_to_word, ScannedDocumentError, inspect_pdf_text_suitability
from services.compress_pdf import compress_pdf_lossless
from services.compress_word import compress_word_document
from werkzeug.datastructures import FileStorage
import io

FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"

class TestDocumentServices(unittest.TestCase):

    def test_ephemeral_workspace_guaranteed_cleanup(self):
        """Verify temporary directories and files are guaranteed purged on exit and exception."""
        ws_path = None
        with EphemeralWorkspace(prefix="test_clean_") as ws:
            ws_path = ws.path
            self.assertTrue(ws_path.exists())
            test_file = ws.secure_path(suffix=".txt")
            test_file.write_text("sensitivedata")
            self.assertTrue(test_file.exists())

        # Outside context: workspace must not exist
        self.assertFalse(ws_path.exists())

        # Test cleanup on exception
        exc_ws_path = None
        try:
            with EphemeralWorkspace(prefix="test_exc_") as ws:
                exc_ws_path = ws.path
                self.assertTrue(exc_ws_path.exists())
                f = ws.secure_path(suffix=".dat")
                f.write_text("data")
                raise RuntimeError("Simulated crash")
        except RuntimeError:
            pass

        self.assertFalse(exc_ws_path.exists())

    def test_sanitize_filename(self):
        """Verify filenames are properly sanitized to prevent directory traversal and header injection."""
        self.assertEqual(sanitize_filename("../../secret.docx"), "secret.docx")
        self.assertEqual(sanitize_filename("valid_doc.docx"), "valid_doc.docx")
        self.assertEqual(sanitize_filename(""), "document")
        self.assertEqual(sanitize_filename("my report (2026) #1.pdf"), "my_report__2026___1.pdf")

    def test_word_to_pdf_conversion(self):
        """Test Word to PDF conversion preserves document structure."""
        complex_docx = FIXTURES_DIR / "complex_document.docx"
        self.assertTrue(complex_docx.exists())

        with EphemeralWorkspace(prefix="test_w2p_") as ws:
            output_pdf = ws.secure_path(suffix=".pdf")
            convert_word_to_pdf(complex_docx, output_pdf, ws)

            self.assertTrue(output_pdf.exists())
            self.assertGreater(output_pdf.stat().st_size, 1000)

            # Inspect converted PDF
            reader = pypdf.PdfReader(str(output_pdf))
            self.assertGreaterEqual(len(reader.pages), 1)
            # Ensure text is selectable in generated PDF
            extracted_text = "".join(page.extract_text() for page in reader.pages)
            self.assertIn("Financial & Technical Summary Report", extracted_text)
            self.assertIn("Multi-Dimensional Performance Metrics", extracted_text)

    def test_pdf_to_word_conversion_text_based(self):
        """Test PDF to Word conversion using pdf2docx for text-based PDFs."""
        text_pdf = FIXTURES_DIR / "text_document.pdf"
        self.assertTrue(text_pdf.exists())

        with EphemeralWorkspace(prefix="test_p2w_") as ws:
            output_docx = ws.secure_path(suffix=".docx")
            meta = convert_pdf_to_word(text_pdf, output_docx)

            self.assertTrue(output_docx.exists())
            self.assertGreater(output_docx.stat().st_size, 1000)
            self.assertGreater(meta["total_chars"], 50)

            # Inspect resulting docx
            doc = docx.Document(str(output_docx))
            full_text = " ".join([p.text for p in doc.paragraphs])
            self.assertTrue("Financial" in full_text or len(doc.tables) > 0)

    def test_pdf_to_word_scanned_detection(self):
        """Verify scanned / image-only PDFs are detected and raise ScannedDocumentError."""
        scanned_pdf = FIXTURES_DIR / "scanned_image_only.pdf"
        self.assertTrue(scanned_pdf.exists())

        with EphemeralWorkspace(prefix="test_scan_") as ws:
            output_docx = ws.secure_path(suffix=".docx")
            with self.assertRaises(ScannedDocumentError) as ctx:
                convert_pdf_to_word(scanned_pdf, output_docx)
            self.assertIn("scanned or image-only", str(ctx.exception).lower())
            self.assertIn("ocr", str(ctx.exception).lower())

    def test_compress_pdf_lossless(self):
        """Test lossless PDF compression reports genuine reductions or optimal state."""
        text_pdf = FIXTURES_DIR / "text_document.pdf"
        self.assertTrue(text_pdf.exists())

        with EphemeralWorkspace(prefix="test_cmp_pdf_") as ws:
            output_pdf = ws.secure_path(suffix=".pdf")
            stats = compress_pdf_lossless(text_pdf, output_pdf, ws)

            self.assertTrue(output_pdf.exists())
            self.assertIn("original_size", stats)
            self.assertIn("compressed_size", stats)
            self.assertIn("percent_saved", stats)
            self.assertIn("is_reduced", stats)
            # Size must never be larger than original
            self.assertLessEqual(stats["compressed_size"], stats["original_size"])

    def test_compress_pdf_strong_mode(self):
        """Test strong mode PDF compression optimizing embedded images."""
        text_pdf = FIXTURES_DIR / "text_document.pdf"
        self.assertTrue(text_pdf.exists())

        with EphemeralWorkspace(prefix="test_cmp_pdf_strong_") as ws:
            output_pdf = ws.secure_path(suffix=".pdf")
            stats = compress_pdf_lossless(text_pdf, output_pdf, ws, strong_mode=True)

            self.assertTrue(output_pdf.exists())
            self.assertTrue(stats["strong_mode"])
            self.assertLessEqual(stats["compressed_size"], stats["original_size"])

    def test_compress_word_safe_and_image_optimization(self):
        """Test safe Word compression and optional embedded image optimization."""
        huge_doc = FIXTURES_DIR / "huge_image_document.docx"
        self.assertTrue(huge_doc.exists())

        # 1. Safe mode (no image downsampling)
        with EphemeralWorkspace(prefix="test_w_safe_") as ws:
            output_safe = ws.secure_path(suffix=".docx")
            stats_safe = compress_word_document(huge_doc, output_safe, ws, optimize_images=False)
            self.assertTrue(output_safe.exists())
            self.assertEqual(stats_safe["images_optimized"], 0)

        # 2. Image optimization mode
        with EphemeralWorkspace(prefix="test_w_img_") as ws:
            output_opt = ws.secure_path(suffix=".docx")
            stats_opt = compress_word_document(huge_doc, output_opt, ws, optimize_images=True)
            self.assertTrue(output_opt.exists())
            self.assertTrue(stats_opt["is_reduced"])
            self.assertGreater(stats_opt["percent_saved"], 0.0)
            self.assertGreater(stats_opt["images_optimized"], 0)

    def test_excel_to_pdf_conversion(self):
        """Test converting an Excel spreadsheet to PDF."""
        xlsx_path = FIXTURES_DIR / "spreadsheet.xlsx"
        self.assertTrue(xlsx_path.exists())

        from services.office_to_pdf import convert_office_to_pdf
        with EphemeralWorkspace(prefix="test_xl2p_") as ws:
            out_pdf = ws.secure_path(suffix=".pdf")
            convert_office_to_pdf(xlsx_path, out_pdf, ws)
            self.assertTrue(out_pdf.exists())
            self.assertGreater(out_pdf.stat().st_size, 1000)
            reader = pypdf.PdfReader(str(out_pdf))
            self.assertGreaterEqual(len(reader.pages), 1)

    def test_powerpoint_to_pdf_conversion(self):
        """Test converting a PowerPoint presentation to PDF."""
        pptx_path = FIXTURES_DIR / "presentation.pptx"
        self.assertTrue(pptx_path.exists())

        from services.office_to_pdf import convert_office_to_pdf
        with EphemeralWorkspace(prefix="test_ppt2p_") as ws:
            out_pdf = ws.secure_path(suffix=".pdf")
            convert_office_to_pdf(pptx_path, out_pdf, ws)
            self.assertTrue(out_pdf.exists())
            self.assertGreater(out_pdf.stat().st_size, 1000)
            reader = pypdf.PdfReader(str(out_pdf))
            self.assertGreaterEqual(len(reader.pages), 1)

    def test_pdf_merge_and_split(self):
        """Test PDF merge and page extraction."""
        from services.pdf_tools import merge_pdf_documents, split_pdf_document, PDFMergeError, PDFSplitError

        pdf1 = FIXTURES_DIR / "text_document.pdf"
        pdf2 = FIXTURES_DIR / "multi_page_sample.pdf"
        self.assertTrue(pdf1.exists())
        self.assertTrue(pdf2.exists())

        # Test Merge
        with EphemeralWorkspace(prefix="test_merge_") as ws:
            merged_pdf = ws.secure_path(suffix=".pdf")
            stats = merge_pdf_documents([pdf1, pdf2], merged_pdf)
            self.assertTrue(merged_pdf.exists())
            self.assertEqual(stats["files_merged"], 2)
            reader = pypdf.PdfReader(str(merged_pdf))
            self.assertGreaterEqual(len(reader.pages), 5)

        # Test Split (Extract pages 2 and 4 from 4-page sample)
        with EphemeralWorkspace(prefix="test_split_") as ws:
            split_pdf = ws.secure_path(suffix=".pdf")
            stats = split_pdf_document(pdf2, split_pdf, "2, 4")
            self.assertTrue(split_pdf.exists())
            self.assertEqual(stats["extracted_pages"], 2)
            self.assertEqual(stats["selected_pages_1based"], [2, 4])
            reader = pypdf.PdfReader(str(split_pdf))
            self.assertEqual(len(reader.pages), 2)

        # Test Split range validation errors
        with EphemeralWorkspace(prefix="test_split_err_") as ws:
            dummy_out = ws.secure_path(suffix=".pdf")
            with self.assertRaises(PDFSplitError):
                split_pdf_document(pdf2, dummy_out, "99")  # out of bounds
            with self.assertRaises(PDFSplitError):
                split_pdf_document(pdf2, dummy_out, "5-2")  # invalid range

    def test_metadata_scrubber(self):
        """Test stripping metadata and tracking properties from PDF and DOCX."""
        from services.metadata_scrubber import sanitize_document_metadata
        import pymupdf

        # 1. Test PDF metadata wipe
        text_pdf = FIXTURES_DIR / "text_document.pdf"
        with EphemeralWorkspace(prefix="test_scrub_pdf_") as ws:
            clean_pdf = ws.secure_path(suffix=".pdf")
            stats = sanitize_document_metadata(text_pdf, clean_pdf, "pdf")
            self.assertTrue(clean_pdf.exists())
            self.assertGreaterEqual(stats["fields_scrubbed"], 1)

            # Check that metadata is wiped in resulting PDF
            doc = pymupdf.open(str(clean_pdf))
            meta = doc.metadata or {}
            self.assertEqual(meta.get("author", ""), "")
            self.assertEqual(meta.get("creator", ""), "")
            self.assertEqual(meta.get("producer", ""), "")
            doc.close()

        # 2. Test DOCX metadata wipe
        meta_docx = FIXTURES_DIR / "doc_with_metadata.docx"
        with EphemeralWorkspace(prefix="test_scrub_docx_") as ws:
            clean_docx = ws.secure_path(suffix=".docx")
            stats = sanitize_document_metadata(meta_docx, clean_docx, "docx")
            self.assertTrue(clean_docx.exists())
            self.assertGreaterEqual(stats["fields_scrubbed"], 1)

            # Check that author and identifying fields were scrubbed
            doc = docx.Document(str(clean_docx))
            self.assertEqual(doc.core_properties.author or "", "")
            self.assertEqual(doc.core_properties.title or "", "")

if __name__ == "__main__":
    unittest.main()

